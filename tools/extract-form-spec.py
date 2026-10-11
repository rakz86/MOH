# Prints the printed request form's design as measured facts.
#
# The approved workbook is the source of truth for that form's layout, not
# the script that first generated it - it was hand-tuned afterwards. Run this
# against the current file and compare with forms/PRINT-FORM-SPEC.md whenever
# either might have moved on.
#
#   python tools/extract-form-spec.py [path-to-workbook]
#
import collections
import os
import sys

import openpyxl
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

DEFAULT = r"D:\MOH\Perio-Instruments-Request Form Oct 2026.xlsx"
PATH = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
if not os.path.exists(PATH):
    sys.exit("Not found: %s  -  pass the workbook path as an argument." % PATH)

wb = openpyxl.load_workbook(PATH)


def argb(c):
    try:
        return str(c.rgb) if c and c.type == "rgb" else None
    except Exception:
        return None


def describe(ws, name):
    print("=" * 72)
    print(name)
    print("=" * 72)

    ps, pm = ws.page_setup, ws.page_margins
    print("PAGE")
    print("  orientation %s   paper %s   fitToPage %s  fitToWidth %s  fitToHeight %s"
          % (ps.orientation, ps.paperSize, ws.sheet_properties.pageSetUpPr.fitToPage,
             ps.fitToWidth, ps.fitToHeight))
    print("  margins L%s R%s T%s B%s  header %s footer %s"
          % (pm.left, pm.right, pm.top, pm.bottom, pm.header, pm.footer))
    print("  print_area %s   title_rows %s   gridlines %s   hCentered %s"
          % (ws.print_area, ws.print_title_rows, ws.sheet_view.showGridLines,
             ws.print_options.horizontalCentered))
    print("  freeze %s   breaks after rows %s"
          % (ws.freeze_panes, sorted(b.id for b in ws.row_breaks.brk)))
    try:
        print("  footer L %r | R %r" % (ws.oddFooter.left.text, ws.oddFooter.right.text))
    except Exception:
        pass
    print("  images %d" % len(ws._images))

    print("COLUMNS  (width units -> px = w*7+5)")
    total = 0
    for i in range(1, 9):
        w = ws.column_dimensions[get_column_letter(i)].width
        px = w * 7 + 5
        total += px
        print("  %s  %-6s %4.0fpx   %s" % (get_column_letter(i), w, px,
                                           ws.cell(row=11, column=i).value or
                                           ws.cell(row=12, column=i).value or ""))
    print("  TOTAL %.0fpx   (A4 portrait usable at these margins = %.0fpx)"
          % (total, (8.2677 - pm.left - pm.right) * 96))

    print("ROW HEIGHTS (pt)")
    for r in range(1, 13):
        v = str(ws.cell(row=r, column=1).value or "")[:46]
        print("  %2d  %5s  %s" % (r, ws.row_dimensions[r].height, v))
    body = collections.Counter(ws.row_dimensions[r].height
                               for r in range(13, ws.max_row + 1))
    print("  body rows: %s" % dict(sorted(body.items(), key=lambda x: -x[1])))

    print("FONTS / FILLS by role")
    seen = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, 9):
            cell = ws.cell(row=r, column=c)
            if cell.value is None and not (cell.fill and cell.fill.fill_type):
                continue
            f, fl = cell.font, cell.fill
            key = (f.name, f.size, bool(f.bold), bool(f.italic), argb(f.color),
                   argb(fl.fgColor) if fl and fl.fill_type == "solid" else None)
            if key not in seen:
                seen[key] = (r, c, str(cell.value)[:34] if cell.value else "")
    print("  %-9s %-5s %-5s %-6s %-10s %-10s  example" %
          ("font", "size", "bold", "italic", "colour", "fill"))
    for k, v in sorted(seen.items(), key=lambda x: x[1]):
        print("  %-9s %-5s %-5s %-6s %-10s %-10s  r%d c%d %s" %
              (k[0], k[1], k[2], k[3], k[4] or "-", k[5] or "-", v[0], v[1], v[2]))
    print()


describe(wb["Per-Clinic Form"], "PER-CLINIC FORM")
describe(wb["Centre-Level Form"], "CENTRE-LEVEL FORM")

ws = wb["Per-Clinic Form"]
print("=" * 72)
print("FIXED TEXT")
print("=" * 72)
for addr in ("A2", "E2", "A3", "A4", "E4", "A5", "E5", "A7", "A8", "A9"):
    print("  %-4s %r" % (addr, ws[addr].value))
print()
print("ROW ARCHETYPES")
for label, r in (("set banner", 13), ("item", 14), ("quota group", 17),
                 ("variant", 18)):
    cells = [ws.cell(row=r, column=c) for c in range(1, 9)]
    print("  %-12s row %-3d h=%-5s fills=%s" % (
        label, r, ws.row_dimensions[r].height,
        [argb(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
         for c in cells]))
    print("               std cell=%r  desc=%r" % (cells[3].value, str(cells[2].value)[:50]))
